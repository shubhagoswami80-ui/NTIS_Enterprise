"""
Fix 42 — Option Chain End-to-End Collector.

First controlled end-to-end benchmark:
    20 symbols
      -> capture the real iCharts request for each symbol
      -> exact direct replay with 429 retry/backoff
      -> validate response
      -> map aaData[3:64] (70-field row) to the 61 browser columns
      -> add Symbol / Expiry / ATM / ATM_Flag
      -> ONE XLSX with Data / Status / Run

This intentionally uses browser pages only for request-contract acquisition.
The actual response used for the report is the direct replay response.

No scheduler. No 220-symbol run unless the user later requests the higher
count. Output is written only after the complete benchmark finishes.
"""

from __future__ import annotations

import asyncio
import json
import random
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

OPTION_CHAIN = "https://www.icharts.in/opt/OptionChain.php"
TABLE_MARKER = "OptionChainTable_Beta_v19.php"

HEADERS_61 = [
    "Calls","Puts","Buildup","Trend","Time","Vega","Theta","Gamma","Delta",
    "IV Chg%","OI Chg%","OI Chg","OI","Volume Chg %","Volume Chg","Volume",
    "OH/OL","Open","High","Low","Close Price Chg %","Close Price Chg",
    "Prev Close Price","Close Price","LTP Chg %","LTP Chg","LTP","VWAP",
    "Bid","Strike Price","PE-CE OI","PE-CE OI Chg","VWAP","LTP","LTP Chg",
    "LTP Chg %","Close Price","Prev Close Price","Close Price Chg",
    "Close Price Chg %","Open","High","Low","OH/OL","Volume","Volume Chg",
    "Volume Chg %","OI Chg","OI Chg%","IV","IV Chg%","Delta","Gamma","Theta",
    "Vega","PCR-OI","PCR-OI Chg","PCR-Vol","Time","Trend","Buildup"
]

DATA_HEADERS = ["Symbol", "Expiry", "ATM", "ATM_Flag"] + HEADERS_61

def _form(body):
    p = parse_qs(body or "", keep_blank_values=True)
    return {k: (v[0] if v else "") for k, v in p.items()}

def _signature(data):
    if not isinstance(data, dict):
        return None

    rows = data.get("aaData") or []
    rows_list = rows if isinstance(rows, list) else []

    # Strike Price is response row index 32 in the validated 70-field
    # iCharts payload. Keep the list even for NO_DATA (empty list).
    strikes = [
        str(row[32]).strip()
        for row in rows_list
        if isinstance(row, list) and len(row) > 32
    ]

    return {
        "server_symbol": str(data.get("symbol_1", "")),
        "expiry": str(data.get("optExpDate", "")),
        "atm": str(data.get("strikePriceATM", "")),
        "optStrike": str(data.get("optStrike", "")),
        "rows": len(rows_list),
        "strikes": strikes,
        "rows_data": rows_list,
    }

async def _capture_one(context, symbol, timeout_seconds):
    started = time.perf_counter()
    page = await context.new_page()

    def is_target(req):
        if TABLE_MARKER not in req.url or req.method.upper() != "POST":
            return False
        params = _form(req.post_data or "")
        return params.get("optSymbol", "").upper() == symbol.upper()

    try:
        await page.goto(
            OPTION_CHAIN,
            wait_until="domcontentloaded",
            timeout=timeout_seconds * 1000,
        )
        await page.wait_for_timeout(1000)

        options = await page.locator("#optSymbol option").evaluate_all(
            """els => els.map(e => String(e.value || '').toUpperCase())"""
        )
        if symbol.upper() not in options:
            raise RuntimeError("Symbol is not present in iCharts #optSymbol.")

        async with page.expect_request(
            is_target, timeout=timeout_seconds * 1000
        ) as req_info:
            await page.select_option("#optSymbol", symbol)

        req = await req_info.value
        response = await req.response()
        if response is None:
            raise RuntimeError("Captured request has no response.")

        body = req.post_data or ""
        response_text = await response.text()
        captured_json = json.loads(response_text)
        sig = _signature(captured_json)

        if sig is None:
            raise RuntimeError("Captured response is not a JSON object.")

        headers = {
            k: v
            for k, v in req.headers.items()
            if k.lower() in {
                "content-type", "referer", "origin",
                "x-requested-with", "accept"
            }
        }
        headers.pop("content-length", None)

        return {
            "symbol": symbol,
            "request_url": req.url,
            "request_body": body,
            "request_params": _form(body),
            "request_headers": headers,
            "captured_http": response.status,
            "captured_signature": sig,
            "capture_ms": round((time.perf_counter() - started) * 1000, 1),
        }
    finally:
        await page.close()

async def _replay_one(context, captured, timeout_seconds, max_retries,
                      backoff_initial_ms, backoff_max_ms,
                      request_gap_ms, jitter_ms):
    symbol = captured["symbol"]
    last_error = ""
    started = time.perf_counter()

    # Small controlled pacing between direct requests.
    if request_gap_ms:
        await asyncio.sleep(request_gap_ms / 1000.0)

    for attempt in range(1, max_retries + 2):
        try:
            headers = dict(captured["request_headers"])
            response = await context.request.post(
                captured["request_url"],
                data=captured["request_body"],
                headers=headers,
                timeout=timeout_seconds * 1000,
            )
            text = await response.text()

            if response.status == 429:
                retry_after = response.headers.get("retry-after", "")
                try:
                    wait_s = float(retry_after)
                except Exception:
                    wait_s = min(
                        backoff_max_ms,
                        backoff_initial_ms * (2 ** (attempt - 1))
                    ) / 1000.0
                wait_s += random.uniform(0, jitter_ms) / 1000.0
                if attempt <= max_retries:
                    await asyncio.sleep(wait_s)
                    continue
                return {
                    "replay_http": 429,
                    "replay_signature": None,
                    "replay_ms": round((time.perf_counter()-started)*1000,1),
                    "attempts": attempt,
                    "error": "HTTP 429 after retries",
                }

            try:
                data = json.loads(text)
            except Exception as exc:
                return {
                    "replay_http": response.status,
                    "replay_signature": None,
                    "replay_ms": round((time.perf_counter()-started)*1000,1),
                    "attempts": attempt,
                    "error": f"Non-JSON replay response: {exc}; body={text[:200]!r}",
                }

            return {
                "replay_http": response.status,
                "replay_signature": _signature(data),
                "replay_ms": round((time.perf_counter()-started)*1000,1),
                "attempts": attempt,
                "error": "",
            }

        except Exception as exc:
            last_error = str(exc)
            if attempt <= max_retries:
                wait_s = min(
                    backoff_max_ms,
                    backoff_initial_ms * (2 ** (attempt - 1))
                ) / 1000.0
                wait_s += random.uniform(0, jitter_ms) / 1000.0
                await asyncio.sleep(wait_s)
                continue

    return {
        "replay_http": None,
        "replay_signature": None,
        "replay_ms": round((time.perf_counter()-started)*1000,1),
        "attempts": max_retries + 1,
        "error": last_error or "Replay failed",
    }

def _map_rows(symbol, sig):
    rows = sig["rows_data"]
    output = []
    atm = str(sig["atm"]).strip()

    for row in rows:
        if not isinstance(row, list):
            continue
        # Validated production mapping: 70 response fields -> browser's
        # 61 displayed columns by removing the first 3 and last 6 fields.
        mapped = list(row[3:64])
        if len(mapped) != 61:
            continue

        strike = str(mapped[29]).strip()
        output.append(
            [symbol, sig["expiry"], sig["atm"], strike == atm] + mapped
        )
    return output

def _write_workbook(output_root, rows, statuses, run_meta):
    now = datetime.now()
    root = Path(output_root)
    folder = root / str(now.year) / now.strftime("%B").lower() / now.strftime("%Y-%m-%d")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"OptionChain_{now:%Y-%m-%d_%H-%M}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(DATA_HEADERS)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    status_ws = wb.create_sheet("Status")
    status_headers = [
        "Symbol","Captured HTTP","Replay HTTP","Expiry","ATM","optStrike",
        "Rows","Attempts","Capture ms","Replay ms","Status","Error"
    ]
    status_ws.append(status_headers)
    for s in statuses:
        status_ws.append([
            s.get("symbol"), s.get("captured_http"), s.get("replay_http"),
            s.get("expiry"), s.get("atm"), s.get("optStrike"),
            s.get("rows"), s.get("attempts"), s.get("capture_ms"),
            s.get("replay_ms"), s.get("status"), s.get("error","")
        ])

    run_ws = wb.create_sheet("Run")
    run_ws.append(["Field","Value"])
    for k,v in run_meta.items():
        run_ws.append([k,v])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for col in range(1, sheet.max_column + 1):
            width = min(
                28,
                max(10, len(str(sheet.cell(1, col).value or "")) + 2)
            )
            sheet.column_dimensions[get_column_letter(col)].width = width

    wb.save(path)
    return str(path)

async def run_optionchain_end_to_end_collector_v3(
    context, symbols, output_root,
    replay_concurrency=5,
    capture_concurrency=5,
    timeout_seconds=30,
    max_retries=2,
    backoff_initial_ms=1500,
    backoff_max_ms=15000,
    request_gap_ms=100,
    jitter_ms=100,
):
    started = time.perf_counter()

    # Phase 1: acquire exact request contracts with bounded browser-page
    # concurrency. Each symbol still gets its own dedicated page; there is no
    # shared page-state mutation between symbols.
    capture_sem = asyncio.Semaphore(max(1, int(capture_concurrency)))

    async def capture_worker(symbol):
        async with capture_sem:
            try:
                return ("CAPTURE", await _capture_one(
                    context, symbol, timeout_seconds
                ))
            except Exception as exc:
                return ("FAIL", {
                    "symbol": symbol,
                    "capture_ms": 0,
                    "captured_http": None,
                    "error": str(exc),
                })

    capture_results = await asyncio.gather(
        *(capture_worker(symbol) for symbol in symbols)
    )

    captures = [
        payload for kind, payload in capture_results if kind == "CAPTURE"
    ]
    capture_failures = [
        payload for kind, payload in capture_results if kind == "FAIL"
    ]

    sem = asyncio.Semaphore(max(1, int(replay_concurrency)))

    async def replay_worker(c):
        async with sem:
            replay = await _replay_one(
                context, c, timeout_seconds, max_retries,
                backoff_initial_ms, backoff_max_ms,
                request_gap_ms, jitter_ms
            )
            return c, replay

    replay_pairs = await asyncio.gather(
        *(replay_worker(c) for c in captures)
    )

    data_rows = []
    statuses = []
    passed = 0
    no_data = 0
    failed = len(capture_failures)

    for failure in capture_failures:
        statuses.append({
            **failure,
            "replay_http": None,
            "expiry": "",
            "atm": "",
            "optStrike": "",
            "rows": 0,
            "attempts": 0,
            "replay_ms": 0,
            "status": "FAILED",
        })

    for captured, replay in replay_pairs:
        cs = captured["captured_signature"]
        rs = replay.get("replay_signature")
        status = "FAILED"

        if replay.get("replay_http") == 200 and rs:
            if rs["rows"] == 0:
                status = "NO_DATA"
                no_data += 1
            elif (
                rs["server_symbol"].upper() == captured["symbol"].upper()
                and rs["atm"] == cs["atm"]
                and rs.get("strikes", []) == cs.get("strikes", [])
            ):
                status = "PASS"
                passed += 1
                data_rows.extend(_map_rows(captured["symbol"], rs))
            else:
                status = "FAILED"
                failed += 1

        else:
            failed += 1

        statuses.append({
            "symbol": captured["symbol"],
            "captured_http": captured["captured_http"],
            "replay_http": replay.get("replay_http"),
            "expiry": cs["expiry"],
            "atm": cs["atm"],
            "optStrike": cs["optStrike"],
            "rows": rs["rows"] if rs else 0,
            "attempts": replay.get("attempts", 0),
            "capture_ms": captured["capture_ms"],
            "replay_ms": replay.get("replay_ms", 0),
            "status": status,
            "error": (
                replay.get("error", "")
                or (
                    f"Payload mismatch: captured_strikes={cs.get('strikes', [])}; "
                    f"replay_strikes={rs.get('strikes', [])}"
                    if status == "FAILED" and rs is not None
                    else ""
                )
            ),
        })

    # Explicit completeness: only write the final workbook after all requested
    # symbols have been classified.
    overall = (
        "COMPLETE" if passed + no_data == len(symbols) and failed == 0
        else "PARTIAL"
    )

    run_meta = {
        "Report": "Option Chain",
        "Run timestamp": datetime.now().isoformat(timespec="seconds"),
        "Requested symbols": len(symbols),
        "Passed symbols": passed,
        "No-data symbols": no_data,
        "Failed symbols": failed,
        "Data rows": len(data_rows),
        "Capture concurrency": capture_concurrency,
        "Replay concurrency": replay_concurrency,
        "Max retries": max_retries,
        "Backoff initial ms": backoff_initial_ms,
        "Backoff max ms": backoff_max_ms,
        "Request gap ms": request_gap_ms,
        "Jitter ms": jitter_ms,
        "Overall status": overall,
    }

    output_file = _write_workbook(output_root, data_rows, statuses, run_meta)

    return {
        "test": "Option Chain End-to-End Collector",
        "requested": len(symbols),
        "passed": passed,
        "no_data": no_data,
        "failed": failed,
        "data_rows": len(data_rows),
        "elapsed_ms": round((time.perf_counter()-started)*1000, 1),
        "output_file": output_file,
        "results": statuses,
        "overall_status": overall,
    }
