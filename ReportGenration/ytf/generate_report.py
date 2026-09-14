import os
import zipfile
from datetime import datetime
import pandas as pd
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import PatternFill, Font

LOG_FOLDER = r"D:\My-data\Share_P&L\ytdata"

def load_target_data(formatted_date):
    """Load the newest complete day-wise interval snapshot for the target date.

    The day-wise snapshot is preferred. The legacy flat CSV/ZIP remains a
    compatibility fallback during migration.
    """
    day_folder = os.path.join(LOG_FOLDER, formatted_date)
    snapshots = []
    if os.path.isdir(day_folder):
        snapshots = [
            os.path.join(day_folder, name)
            for name in os.listdir(day_folder)
            if name.startswith("snapshot_") and name.endswith(".csv")
            and not name.endswith(".tmp")
        ]
    if snapshots:
        snapshot_path = max(snapshots, key=os.path.getmtime)
        print(f"📖 Loading latest day-wise snapshot: {os.path.basename(snapshot_path)}")
        return pd.read_csv(snapshot_path)

    csv_file = f"nse_log_{formatted_date}.csv"
    zip_file = f"nse_log_{formatted_date}.zip"
    csv_path = os.path.join(LOG_FOLDER, csv_file)
    zip_path = os.path.join(LOG_FOLDER, zip_file)
    if os.path.exists(csv_path):
        print("📖 Loading legacy raw log as compatibility fallback...")
        return pd.read_csv(csv_path, on_bad_lines="warn")
    if os.path.exists(zip_path):
        print("📖 Loading legacy compressed log as compatibility fallback...")
        with zipfile.ZipFile(zip_path, "r") as archive:
            with archive.open(csv_file) as f:
                return pd.read_csv(f, on_bad_lines="warn")
    return None

def generate_excel_report():
    target_date = input("\n📅 Enter target date (YYYY-MM-DD) or press ENTER for today: ").strip()
    if not target_date:
        target_date = datetime.now().strftime("%Y-%m-%d")

    clean_date = target_date.replace("-", "").replace("/", "")
    formatted_date = f"{clean_date[:4]}-{clean_date[4:6]}-{clean_date[6:]}"

    raw_df = load_target_data(formatted_date)

    if raw_df is None or raw_df.empty:
        print(f"❌ Error: No matching logging structures found for: {formatted_date}")
        return

    print(f"⚙ Processing parameters and calculating skew ratios...")
    processed_df = raw_df.copy()

    # Apply Screenshot Bias Skew formula: ((Call_chg - Put_chg) / (Call_chg + Put_chg)) * 100
    processed_df["Bias_%"] = processed_df.apply(
        lambda row: (
            round(((row["ce_volume_change"] - row["pe_volume_change"]) / (row["ce_volume_change"] + row["pe_volume_change"])) * 100, 2)
            if (row["ce_volume_change"] + row["pe_volume_change"]) > 0 else 0.0
        ), axis=1
    )

    processed_df["Skew"] = processed_df["Bias_%"].apply(
        lambda x: "Bullish Skew" if x > 0 else ("Bearish Skew" if x < 0 else "Neutral")
    )

    # 15% and 20% Threshold Break logic tracking from your image rules
    def determine_signal(bias):
        abs_bias = abs(bias)
        if abs_bias >= 20:
            return "BROKE (confirmed)"
        elif abs_bias >= 15:
            return "BROKE (immediate)"
        return "Normal Straddle"

    processed_df["Event_Status"] = processed_df["Bias_%"].apply(determine_signal)

    output_filename = f"NSE_Processed_Report_{formatted_date}.xlsx"

    # ==============================================================================
    # 📊 EXCEL ARRAYS INTERACTIVE MATRIX GENERATION AREA
    # ==============================================================================
    with pd.ExcelWriter(output_filename, engine="openpyxl") as writer:
        processed_df.to_excel(writer, sheet_name="Processed_Data", index=False)
        raw_df.to_excel(writer, sheet_name="Raw_Logs", index=False)

        workbook = writer.book
        dashboard = workbook.create_sheet(title="Interactive_Dashboard", index=0)

        # Style Layout Elements
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Slate Blue Header
        
        # Format Dashboard Panel layout
        dashboard["A1"] = "🎯 SIGNAL STRADDLE DROP-DOWN FILTER"
        dashboard["A1"].font = Font(name="Arial", size=14, bold=True, color="1F4E78")
        
        dashboard["A3"] = "Select Signal Status Filter 👉"
        dashboard["A3"].font = Font(name="Arial", size=11, bold=True)
        dashboard["B3"] = "BROKE (confirmed)"  # Seed default item
        dashboard["B3"].font = Font(name="Arial", size=11, bold=True, color="FF0000")

        # Setup Table Headers
        headers = ["Timestamp", "Ticker Symbol", "Expiry", "CE Vol Change", "PE Vol Change", "Bias %", "Event Status"]
        for col_idx, text in enumerate(headers, start=1):
            cell = dashboard.cell(row=5, column=col_idx, value=text)
            cell.font = header_font
            cell.fill = header_fill

        # Map verification dropdown list onto cell B3
        dv = DataValidation(type="list", formula1='"BROKE (confirmed),BROKE (immediate),Normal Straddle"', allow_blank=True)
        dashboard.add_data_validation(dv)
        dv.add(dashboard["B3"])

        # 👑 FIX: Changed lookup column inside index formulas from 'M' to 'N' to match column allocation
        for i in range(6, 150):
            dashboard[f"A{i}"] = f'=IFERROR(INDEX(Processed_Data!A:A, SMALL(IF(Processed_Data!$N:$N=$B$3, ROW(Processed_Data!$N:$N)), ROW({f"A{i-5}"}))), "")'
            dashboard[f"B{i}"] = f'=IFERROR(INDEX(Processed_Data!B:B, SMALL(IF(Processed_Data!$N:$N=$B$3, ROW(Processed_Data!$N:$N)), ROW({f"A{i-5}"}))), "")'
            dashboard[f"C{i}"] = f'=IFERROR(INDEX(Processed_Data!C:C, SMALL(IF(Processed_Data!$N:$N=$B$3, ROW(Processed_Data!$N:$N)), ROW({f"A{i-5}"}))), "")'
            dashboard[f"D{i}"] = f'=IFERROR(INDEX(Processed_Data!E:E, SMALL(IF(Processed_Data!$N:$N=$B$3, ROW(Processed_Data!$N:$N)), ROW({f"A{i-5}"}))), "")'
            dashboard[f"E{i}"] = f'=IFERROR(INDEX(Processed_Data!G:G, SMALL(IF(Processed_Data!$N:$N=$B$3, ROW(Processed_Data!$N:$N)), ROW({f"A{i-5}"}))), "")'
            dashboard[f"F{i}"] = f'=IFERROR(INDEX(Processed_Data!L:L, SMALL(IF(Processed_Data!$N:$N=$B$3, ROW(Processed_Data!$N:$N)), ROW({f"A{i-5}"}))), "")'
            dashboard[f"G{i}"] = f'=IFERROR(INDEX(Processed_Data!N:N, SMALL(IF(Processed_Data!$N:$N=$B$3, ROW(Processed_Data!$N:$N)), ROW({f"A{i-5}"}))), "")'

        # 🎨 ADD COLOR CODING (Conditional Highlights)
        # Green Fill for Bullish Skews, Soft Red/Orange for Bearish Skews
        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        
        # Apply conditional formatting formulas inside openpyxl workspace sheets
        from openpyxl.formatting.rule import CellIsRule
        dashboard.conditional_formatting.add("F6:F150", CellIsRule(operator="greaterThan", formula=["0"], stopIfTrue=True, fill=green_fill))
        dashboard.conditional_formatting.add("F6:F150", CellIsRule(operator="lessThan", formula=["0"], stopIfTrue=True, fill=red_fill))

        # Auto-adjust column sizes for beautiful scannability layouts
        for col in dashboard.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            dashboard.column_dimensions[col_letter].width = max(max_len + 3, 14)

    print(f"🎉 Fully Interactive Excel Report Workbook compiled successfully!")
    print(f"👉 File generated inside your local directory: {output_filename}")

if __name__ == "__main__":
    generate_excel_report()
