import concurrent.futures
import csv
import json
import os
import random
import time
import zipfile
from datetime import datetime
import pandas as pd
import requests
import yfinance as yf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# Config paths matching your storage layout
LOG_FOLDER = r"D:\My-data\Share_P&L\ytdata"
LIVE_EXCEL_PATH = os.path.join(LOG_FOLDER, "NSE_Live_Straddle_Dashboard.xlsx")

# TEST_MODE Switch: Change to False for live sessions
TEST_MODE = True
STATE_CACHE = {}
STATE_CACHE_FILE = os.path.join(LOG_FOLDER, "collector_state.json")
SNAPSHOT_HEADERS = [
    "timestamp", "ticker", "expiry", "total_ce_volume",
    "ce_volume_change", "total_pe_volume", "pe_volume_change",
    "total_ce_oi", "ce_oi_change", "total_pe_oi",
    "pe_oi_change", "Bias_%"
]


def is_market_hours():
    if TEST_MODE:
        return True
    now = datetime.now()
    if now.weekday() >= 5:
        return False
    return (
        now.replace(hour=9, minute=15, second=0)
        <= now
        <= now.replace(hour=15, minute=30, second=0)
    )


def archive_daily_log(target_date_str):
    csv_filename = f"nse_log_{target_date_str}.csv"
    csv_path = os.path.join(LOG_FOLDER, csv_filename)
    zip_path = os.path.join(LOG_FOLDER, f"nse_log_{target_date_str}.zip")
    if os.path.exists(csv_path) and not os.path.exists(zip_path):
        try:
            with zipfile.ZipFile(
                zip_path, "w", zipfile.ZIP_DEFLATED, compresslevel=9
            ) as zipf:
                zipf.write(csv_path, arcname=csv_filename)
            os.remove(csv_path)
            print("📦 Daily uncompressed track cleared.")
        except Exception as e:
            print(f"❌ Zip fail: {e}")


def get_current_fno_tickers():
    try:
        url = "https://nseindia.com"
        response = requests.get(
            url, headers={"User-Agent": "Mozilla/5.0"}, timeout=10
        )
        if response.status_code == 200:
            return sorted(
                list(
                    set(
                        [
                            f"{line.split(',')[1].strip()}.NS"
                            for line in response.text.split("\n")[5:]
                            if len(line.split(",")) > 1
                            and line.split(",")[1].strip()
                            and not line.split(",")[1]
                            .strip()
                            .startswith("MKT")
                        ]
                    )
                )
            )
    except:
        pass
    return ["RELIANCE.NS", "TCS.NS", "INFY.NS", "HDFCBANK.NS", "SBIN.NS"]


def load_persistent_state():
    global STATE_CACHE
    try:
        if os.path.isfile(STATE_CACHE_FILE):
            with open(STATE_CACHE_FILE, "r", encoding="utf-8") as f:
                raw = json.load(f)
            STATE_CACHE = {k: tuple(v) for k, v in raw.items()}
    except Exception as exc:
        print(f"State cache load warning: {exc}")
        STATE_CACHE = {}


def save_persistent_state():
    os.makedirs(LOG_FOLDER, exist_ok=True)
    temp_path = STATE_CACHE_FILE + ".tmp"
    with open(temp_path, "w", encoding="utf-8") as f:
        json.dump(STATE_CACHE, f, separators=(",", ":"))
    os.replace(temp_path, STATE_CACHE_FILE)


def process_and_stream_ticker(ticker_symbol, timestamp_str):
    global STATE_CACHE
    try:
        stock = yf.Ticker(ticker_symbol)
        expirations = stock.options
        total_ce_volume = total_pe_volume = total_ce_oi = total_pe_oi = 0
        nearest_expiry = "N/A"

        if expirations:
            nearest_expiry = expirations[0]
            opt_chain = stock.option_chain(nearest_expiry)
            if opt_chain.calls is not None and not opt_chain.calls.empty:
                total_ce_volume = int(opt_chain.calls["volume"].fillna(0).sum())
                total_ce_oi = int(opt_chain.calls["openInterest"].fillna(0).sum())
            if opt_chain.puts is not None and not opt_chain.puts.empty:
                total_pe_volume = int(opt_chain.puts["volume"].fillna(0).sum())
                total_pe_oi = int(opt_chain.puts["openInterest"].fillna(0).sum())

        if TEST_MODE:
            nearest_expiry = nearest_expiry if nearest_expiry != "N/A" else "2026-09-24"
            total_ce_volume = random.randint(50000, 150000)
            total_pe_volume = random.randint(50000, 150000)
            total_ce_oi = random.randint(10000, 50000)
            total_pe_oi = random.randint(10000, 50000)

        ce_vol_chg = pe_vol_chg = ce_oi_chg = pe_oi_chg = 0
        previous = STATE_CACHE.get(ticker_symbol)
        if previous:
            prev_cv, prev_pv, prev_co, prev_po = previous
            ce_vol_chg = max(0, total_ce_volume - prev_cv)
            pe_vol_chg = max(0, total_pe_volume - prev_pv)
            ce_oi_chg = total_ce_oi - prev_co
            pe_oi_chg = total_pe_oi - prev_po

        STATE_CACHE[ticker_symbol] = (total_ce_volume, total_pe_volume, total_ce_oi, total_pe_oi)
        total_volume_change = ce_vol_chg + pe_vol_chg
        bias = round(((ce_vol_chg - pe_vol_chg) / total_volume_change) * 100, 2) if total_volume_change else 0.0
        return [timestamp_str, ticker_symbol, nearest_expiry, total_ce_volume, ce_vol_chg,
                total_pe_volume, pe_vol_chg, total_ce_oi, ce_oi_chg, total_pe_oi, pe_oi_chg, bias]
    except Exception:
        return None


def write_snapshot_atomic(rows, timestamp):
    if not rows:
        return None
    date_str = timestamp.strftime("%Y-%m-%d")
    time_str = timestamp.strftime("%H%M")
    day_folder = os.path.join(LOG_FOLDER, date_str)
    os.makedirs(day_folder, exist_ok=True)
    final_path = os.path.join(day_folder, f"snapshot_{time_str}.csv")
    if os.path.isfile(final_path) and os.path.getsize(final_path) > 0:
        return final_path
    temp_path = final_path + ".tmp"
    with open(temp_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(SNAPSHOT_HEADERS)
        writer.writerows(rows)
        f.flush()
        os.fsync(f.fileno())
    os.replace(temp_path, final_path)
    return final_path


def apply_modern_styles(
    ws, start_row, title_text, headers, rows_data, header_bg="1F4E78"
):
    title_font = Font(name="Segoe UI", size=11, bold=True, color="1F4E78")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="333333")
    fill_header = PatternFill(
        start_color=header_bg, end_color=header_bg, fill_type="solid"
    )
    fill_zebra = PatternFill(
        start_color="F9FBFD", end_color="F9FBFD", fill_type="solid"
    )
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    ws.cell(row=start_row, column=1, value=title_text).font = title_font
    current_row = start_row + 1

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=h_text)
        cell.font = header_font
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
    current_row += 1

    if not rows_data:
        ws.cell(row=current_row, column=1, value="No current alerts.").font = (
            Font(name="Segoe UI", size=10, italic=True, color="888888")
        )
        current_row += 1
    else:
        for r_idx, r_data in enumerate(rows_data):
            for c_idx, val in enumerate(r_data, start=1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
                if r_idx % 2 == 1:
                    cell.fill = fill_zebra
                if c_idx == 6:
                    b_val = float(val)
                    if b_val > 0:
                        cell.fill = PatternFill(
                            start_color="E2EFDA",
                            end_color="E2EFDA",
                            fill_type="solid",
                        )
                    elif b_val < 0:
                        cell.fill = PatternFill(
                            start_color="FCE4D6",
                            end_color="FCE4D6",
                            fill_type="solid",
                        )
            current_row += 1
    return current_row + 2


def run_analytics_cycle():
    if not is_market_hours():
        return
    cycle_time = datetime.now().replace(second=0, microsecond=0)
    # Keep one shared timestamp for every ticker in this collection cycle.
    now_str = cycle_time.strftime("%Y-%m-%d %H:%M:%S")
    tickers = get_current_fno_tickers()
    print(f"[{now_str}] Executing clean background stream...")
    with concurrent.futures.ThreadPoolExecutor(max_workers=20) as executor:
        futures = [executor.submit(process_and_stream_ticker, ticker, now_str) for ticker in tickers]
        cycle_results = [result for result in (future.result() for future in concurrent.futures.as_completed(futures)) if result is not None]
    snapshot_path = write_snapshot_atomic(cycle_results, cycle_time)
    if snapshot_path:
        save_persistent_state()
        print(f"[{now_str}] Cycle complete: {len(cycle_results)} rows. Snapshot: {snapshot_path}")
    else:
        print(f"[{now_str}] Cycle produced no usable rows.")


if __name__ == "__main__":
    os.makedirs(LOG_FOLDER, exist_ok=True)
    load_persistent_state()
    print(f"NSE option-chain collector started. TEST_MODE= {TEST_MODE}")
    while True:
        try:
            run_analytics_cycle()
        except Exception as exc:
            print(f"Collector cycle error: {exc}")
        time.sleep(300)
