from __future__ import annotations
import json, re, time
from datetime import datetime
from pathlib import Path
from typing import Iterable
import openpyxl

PREFIX="Daywise_Price_and_OI_Summary_"
TS_RE=re.compile(r"_(\d{8})_(\d{6})\.xlsx$", re.I)

def source_root() -> Path:
    return Path(__import__("os").environ.get(
        "NTIS_W73_SOURCE_ROOT",
        r"D:\My-data\Share_P&L\Ichart Data\Screenshot"
    ))

def discover_day_folder(trading_date: str, month="September26") -> Path:
    return source_root()/month/trading_date

def timestamp_from_name(name: str) -> datetime|None:
    m=TS_RE.search(name)
    if not m: return None
    return datetime.strptime(m.group(1)+m.group(2), "%Y%m%d%H%M%S")

def discover_files(trading_date: str, month="September26", cutoff: datetime|None=None) -> list[tuple[Path,datetime]]:
    folder=discover_day_folder(trading_date,month)
    if not folder.exists(): return []
    out=[]
    for p in folder.iterdir():
        if not p.is_file() or p.suffix.lower()!=".xlsx" or not p.name.startswith(PREFIX): continue
        ts=timestamp_from_name(p.name)
        if ts is not None and (cutoff is None or ts <= cutoff): out.append((p,ts))
    return sorted(out,key=lambda x:x[1])

def read_workbook(path: Path, ts: datetime) -> list[dict]:
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws=wb[wb.sheetnames[0]]
        rows=ws.iter_rows(values_only=True)
        header=next(rows,None)
        if not header: return []
        header=[str(x).strip() if x is not None else "" for x in header]
        result=[]
        for vals in rows:
            r={header[i]: (vals[i] if i<len(vals) else None) for i in range(len(header)) if header[i]}
            sym=str(r.get("Symbol","")).strip()
            if not sym: continue
            r["_trading_date"]=ts.strftime("%Y-%m-%d")
            r["_observation_timestamp"]=ts.isoformat(timespec="seconds")
            r["_source_file"]=path.name
            result.append(r)
        return result
    finally:
        wb.close()

def ingest(trading_date: str, month="September26", cache_root: str|Path="07_OUTPUT/live_cache") -> dict:
    cache=Path(cache_root); cache.mkdir(parents=True,exist_ok=True)
    state_path=cache/"ingest_state.json"
    state=json.loads(state_path.read_text()) if state_path.exists() else {"files":{}}
    count=0; symbols=set(); intervals=[]
    for p,ts in discover_files(trading_date,month):
        key=p.name
        if key in state["files"]: continue
        try:
            rows=read_workbook(p,ts)
        except Exception:
            continue
        out=cache/f"{trading_date}.jsonl"
        with out.open("a",encoding="utf-8") as f:
            for r in rows:
                f.write(json.dumps(r,ensure_ascii=False,default=str,separators=(",",":"))+"\n")
                count+=1; symbols.add(r["Symbol"].strip().upper())
        state["files"][key]={"timestamp":ts.isoformat(),"rows":len(rows)}
        intervals.append(ts.isoformat())
    state_path.write_text(json.dumps(state,indent=2),encoding="utf-8")
    return {"status":"OK","trading_date":trading_date,"new_rows":count,"new_symbols":len(symbols),
            "new_intervals":intervals,"cache":str(cache/f"{trading_date}.jsonl")}
