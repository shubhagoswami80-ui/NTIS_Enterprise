from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import os
import re

import openpyxl

DEFAULT_SOURCE_ROOT = r"D:\My-data\Share_P&L\Ichart Data\Screenshot"
DEFAULT_FILE_PREFIX = "Daywise_Price_and_OI_Summary_"
_TIMESTAMP_RE = re.compile(r"_(\d{8})_(\d{6})\.xlsx$", re.IGNORECASE)


@dataclass(frozen=True)
class W73SourceConfig:
    source_root: Path
    file_prefix: str = DEFAULT_FILE_PREFIX
    allowed_extensions: tuple[str, ...] = (".xlsx",)

    @classmethod
    def from_environment(cls) -> "W73SourceConfig":
        return cls(
            source_root=Path(os.getenv("NTIS_W73_SOURCE_ROOT", DEFAULT_SOURCE_ROOT)),
            file_prefix=os.getenv("NTIS_W73_SOURCE_FILE_PREFIX", DEFAULT_FILE_PREFIX),
        )


@dataclass(frozen=True)
class SourceFile:
    path: Path
    timestamp: datetime


class W73SourceAdapter:
    """Read-only, non-recursive adapter for the approved Daywise XLSX tree."""

    def __init__(self, config: W73SourceConfig | None = None) -> None:
        self.config = config or W73SourceConfig.from_environment()

    def day_path(self, month_folder: str, trading_date: str) -> Path:
        return self.config.source_root / month_folder / trading_date

    def list_intervals(
        self,
        month_folder: str,
        trading_date: str,
        *,
        cutoff: datetime | None = None,
    ) -> list[SourceFile]:
        folder = self.day_path(month_folder, trading_date)
        if not folder.is_dir():
            raise FileNotFoundError(f"W73 source day folder not found: {folder}")

        result: list[SourceFile] = []
        for path in folder.iterdir():  # deliberately non-recursive
            if not path.is_file() or path.suffix.lower() not in self.config.allowed_extensions:
                continue
            if not path.name.startswith(self.config.file_prefix):
                continue
            match = _TIMESTAMP_RE.search(path.name)
            if not match:
                continue
            timestamp = datetime.strptime(
                match.group(1) + match.group(2), "%Y%m%d%H%M%S"
            )
            if cutoff is not None and timestamp > cutoff:
                continue
            result.append(SourceFile(path, timestamp))

        return sorted(result, key=lambda x: x.timestamp)

    def latest_interval(self, month_folder: str, trading_date: str) -> SourceFile:
        files = self.list_intervals(month_folder, trading_date)
        if not files:
            raise FileNotFoundError(
                f"No approved W73 interval XLSX files found in {self.day_path(month_folder, trading_date)}"
            )
        return files[-1]

    @staticmethod
    def read_workbook(path: Path) -> tuple[list[str], list[dict[str, object]]]:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
            header = list(next(rows))
            records = [dict(zip(header, row)) for row in rows]
            return header, records
        finally:
            wb.close()

    def read_interval(
        self,
        month_folder: str,
        trading_date: str,
        *,
        cutoff: datetime | None = None,
    ) -> tuple[SourceFile, list[str], list[dict[str, object]]]:
        files = self.list_intervals(month_folder, trading_date, cutoff=cutoff)
        if not files:
            raise FileNotFoundError(
                f"No W73 interval exists at or before cutoff in {self.day_path(month_folder, trading_date)}"
            )
        source = files[-1]
        header, records = self.read_workbook(source.path)
        return source, header, records
