from __future__ import annotations
import json, re, sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = Path(__import__("os").environ.get(
    "NTIS_W73_SOURCE_ROOT",
    r"D:\My-data\Share_P&L\Ichart Data\Screenshot"
))
PREFIX = "Daywise_Price_and_OI_Summary_"
TS_RE = re.compile(r"_(\d{8})_(\d{6})\.xlsx$", re.I)

def main(month: str, trading_date: str):
    folder = SOURCE_ROOT / month / trading_date
    files = []
    for p in folder.iterdir():
        if p.is_file() and p.suffix.lower() == ".xlsx" and p.name.startswith(PREFIX):
            m = TS_RE.search(p.name)
            if m:
                ts = datetime.strptime(m.group(1)+m.group(2), "%Y%m%d%H%M%S")
                files.append((p, ts))
    files.sort(key=lambda x:x[1])
    rows = defaultdict(list)
    for p, ts in files:
        wb=openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            ws=wb[wb.sheetnames[0]]
            it=ws.iter_rows(values_only=True)
            h=next(it)
            idx={str(v).strip():i for i,v in enumerate(h) if v is not None}
            required={"Symbol","Open","High","Low","Close"}
            if not required.issubset(idx):
                raise RuntimeError(f"MISSING_FIELDS:{p.name}")
            for vals in it:
                sym=str(vals[idx["Symbol"]] or "").strip().upper()
                if not sym: continue
                def num(k):
                    v=vals[idx[k]]
                    return float(v) if v is not None else None
                rows[sym].append({
                    "timestamp":ts,"open":num("Open"),"high":num("High"),
                    "low":num("Low"),"close":num("Close")
                })
        finally:
            wb.close()

    result=[]
    cumulative_ok=0
    cumulative_exceptions=0
    breakout_symbols=0
    no_break_symbols=0
    for sym, rs in sorted(rows.items()):
        rs.sort(key=lambda r:r["timestamp"])
        t0=rs[0]["timestamp"]
        cutoff=t0+timedelta(minutes=15)
        orb=[r for r in rs if r["timestamp"]<=cutoff]
        later=[r for r in rs if r["timestamp"]>cutoff]
        high_ok=all(orb[i]["high"] >= orb[i-1]["high"] for i in range(1,len(orb)))
        low_ok=all(orb[i]["low"] <= orb[i-1]["low"] for i in range(1,len(orb)))
        open_ok=all(r["open"] == rs[0]["open"] for r in rs if r["open"] is not None)
        hi=max(r["high"] for r in orb)
        lo=min(r["low"] for r in orb)
        up=next((r for r in later if r["close"]>hi),None)
        dn=next((r for r in later if r["close"]<lo),None)
        hit=min([x for x in (up,dn) if x], key=lambda x:x["timestamp"]) if any((up,dn)) else None
        if high_ok and low_ok:
            cumulative_ok += 1
        else:
            cumulative_exceptions += 1
        if hit: breakout_symbols += 1
        else: no_break_symbols += 1
        result.append({
            "symbol":sym,"observations":len(rs),"t0":t0.isoformat(),
            "orb_cutoff":cutoff.isoformat(),"orb_high":hi,"orb_low":lo,
            "high_non_decreasing":high_ok,"low_non_increasing":low_ok,
            "open_constant":open_ok,
            "breakout": None if hit is None else {
                "timestamp":hit["timestamp"].isoformat(),
                "direction":"UP" if up is hit else "DOWN",
                "close":hit["close"],
            }
        })
    summary={
        "status":"PASS_ORB_SOURCE_GATE" if files and rows and cumulative_exceptions==0 else "REVIEW_ORB_SOURCE_GATE",
        "month":month,"trading_date":trading_date,"files":len(files),
        "symbols":len(rows),"cumulative_ohlc_ok":cumulative_ok,
        "cumulative_ohlc_exceptions":cumulative_exceptions,
        "symbols_with_breakout":breakout_symbols,
        "symbols_without_breakout":no_break_symbols,
        "first_timestamp":files[0][1].isoformat() if files else None,
        "orb_definition":"first_available_timestamp + 15 minutes; OHLC rows <= cutoff; breakout rows > cutoff",
        "symbols":result,
    }
    out=ROOT/"07_OUTPUT"/"w73_full_source_orb_gate_v1.json"
    out.write_text(json.dumps(summary,indent=2),encoding="utf-8")
    print(json.dumps({k:v for k,v in summary.items() if k!="symbols"},indent=2))
    print(f"REPORT={out}")
    return 0 if summary["status"]=="PASS_ORB_SOURCE_GATE" else 2

if __name__=="__main__":
    if len(sys.argv)!=3:
        print("usage: python w73_full_source_orb_gate.py September26 2026-09-18")
        raise SystemExit(2)
    raise SystemExit(main(sys.argv[1],sys.argv[2]))

